"""Word-экспорт протокола ТКМ (python-docx)"""
import datetime
from io import BytesIO
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def _set_cell_bg(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _set_row_bold(row, size=10):
    for cell in row.cells:
        for para in cell.paragraphs:
            for run in para.runs:
                run.bold = True
                run.font.size = Pt(size)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


ACTION_HEX = {
    "тонизация":     "1a6b3a",
    "седация":       "8b2020",
    "обезболивание": "7a5500",
}


def generate_word(
    scores: dict,
    protocol: list,
    selected_symptoms: list,
    tcm_text: str,
    filename: str = None
):
    doc = Document()

    # ── Поля страницы ──
    for section in doc.sections:
        section.top_margin    = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    # ── Заголовок ──
    title = doc.add_heading("ПРОТОКОЛ АКУПУНКТУРНЫХ ТОЧЕК  ТКМ", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].font.size = Pt(16)
    title.runs[0].font.color.rgb = RGBColor(0x1a, 0x4a, 0x6a)

    doc.add_paragraph(
        f"Дата: {datetime.date.today().strftime('%d.%m.%Y')}    "
        f"Точек в протоколе: {len(protocol)}"
    ).runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    doc.add_paragraph()

    # ── Жалобы ──
    if selected_symptoms:
        h = doc.add_heading("Жалобы пациента", 2)
        h.runs[0].font.color.rgb = RGBColor(0x1a, 0x3a, 0x5a)
        for s in selected_symptoms:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(s).font.size = Pt(11)

    doc.add_paragraph()

    # ── Таблица протокола ──
    h = doc.add_heading("Протокол акупунктурных точек", 2)
    h.runs[0].font.color.rgb = RGBColor(0x1a, 0x3a, 0x5a)

    if protocol:
        cols = ["№", "Точка", "Меридиан", "Тип точки", "Действие", "Правило"]
        table = doc.add_table(rows=1, cols=len(cols))
        table.style = "Table Grid"

        # Заголовок таблицы
        hdr_cells = table.rows[0].cells
        for i, h_txt in enumerate(cols):
            hdr_cells[i].text = h_txt
            hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell_bg(hdr_cells[i], "1a3a5a")
        _set_row_bold(table.rows[0], size=10)

        # Ширины колонок
        widths = [Cm(1.0), Cm(2.0), Cm(3.5), Cm(3.5), Cm(2.5), Cm(6.5)]
        for i, w in enumerate(widths):
            for row in table.rows:
                row.cells[i].width = w

        # Строки точек
        for idx, p in enumerate(protocol, 1):
            row = table.add_row()
            vals = [
                str(idx),
                p["point"],
                f"{p['name']} ({p['code']})",
                p.get("point_type", "—"),
                p["action"].upper(),
                p["rule"],
            ]
            bg = ACTION_HEX.get(p["action"], "243547")
            for i, val in enumerate(vals):
                cell = row.cells[i]
                cell.text = val
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                if i <= 1:
                    _set_cell_bg(cell, bg)
                    for para in cell.paragraphs:
                        for run in para.runs:
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                            run.bold = True
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(10)

        # Описания точек
        doc.add_paragraph()
        h3 = doc.add_heading("Описание точек", 3)
        h3.runs[0].font.color.rgb = RGBColor(0x2a, 0x5a, 0x3a)
        for p in protocol:
            if p.get("point_desc"):
                para = doc.add_paragraph()
                r = para.add_run(f"{p['point']} [{p.get('point_type', '')}]:  ")
                r.bold = True
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(0x1a, 0x6b, 0x3a)
                para.add_run(p["point_desc"]).font.size = Pt(10)

    # ── Принципы ТКМ ──
    doc.add_paragraph()
    doc.add_page_break()
    h2 = doc.add_heading("Принципы ТКМ — Анализ состояния", 2)
    h2.runs[0].font.color.rgb = RGBColor(0x1a, 0x3a, 0x5a)

    for line in tcm_text.split("\n"):
        if line.startswith("═") or line.startswith("─"):
            continue
        p = doc.add_paragraph()
        if line.startswith("  ▸") or line.startswith("  •"):
            p.style = "List Bullet"
            p.add_run(line.strip()).font.size = Pt(10)
        elif line.isupper() and len(line) > 3:
            r = p.add_run(line)
            r.bold = True
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(0x2a, 0x5a, 0x7a)
        else:
            p.add_run(line).font.size = Pt(10)

    if filename:
        doc.save(filename)
        return None
    else:
        buf = BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf
